"""One-off: dump formulas + cached values of plan_financier.xlsx (not shipped)."""
import sys
import openpyxl

PATH = "docs/plan_financier.xlsx"
wb_f = openpyxl.load_workbook(PATH, data_only=False)
wb_v = openpyxl.load_workbook(PATH, data_only=True)

sheets = sys.argv[1:] or wb_f.sheetnames
max_rows = 60
for name in sheets:
    ws_f, ws_v = wb_f[name], wb_v[name]
    print(f"\n===== {name} dims={ws_f.dimensions} max_row={ws_f.max_row} max_col={ws_f.max_column}")
    for row in ws_f.iter_rows(min_row=1, max_row=min(ws_f.max_row, max_rows)):
        for c in row:
            if c.value is None:
                continue
            v = ws_v[c.coordinate].value
            if isinstance(c.value, str) and c.value.startswith("="):
                print(f"{c.coordinate}: {c.value!r}  => {v!r}")
            else:
                print(f"{c.coordinate}: {c.value!r}")
