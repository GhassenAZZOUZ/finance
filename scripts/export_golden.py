"""
One-off (Phase 0, not shipped): export golden test data to tests/fixtures/golden.json.

Expected values are the CACHED VALUES computed by the spreadsheet itself (desktop Excel
recalculation of each scenario, see recalc_scenarios.ps1) — never values from our own code.
reference_engine.py is only used as a cross-check (the export aborts on any mismatch).

Usage:
    powershell -File scripts/recalc_scenarios.ps1 -OutDir <dir>
    python scripts/export_golden.py <dir>
"""
from __future__ import annotations

import datetime as dt
import json
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
import reference_engine as ref  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "tests", "fixtures", "golden.json")
DIGITS = 9  # plenty for a 0.01 EUR tolerance, keeps the file small


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        r = round(float(v), DIGITS)
        return 0.0 if r == 0 else r
    return v


def val(v):
    """Excel cell value -> JSON value (dates as 'YYYY-MM' months, text kept verbatim)."""
    if isinstance(v, dt.datetime):
        return month_of(v)
    return num(v)


def month_of(v):
    return v.strftime("%Y-%m") if isinstance(v, dt.datetime) else None


def lines(ws, rows):
    return [{"label": ws[f"A{r}"].value, "amount": num(ws[f"B{r}"].value) or 0.0} for r in rows]


def export_scenario(path, meta):
    wv = openpyxl.load_workbook(path, data_only=True)
    b, c, p, k, s, su = (wv[x] for x in ("Budget", "Crédits", "Plan", "Calcul", "Synthèse", "Suivi réel"))

    # cross-check with the independent reference implementation
    checked, errors, _ = ref.compare(path)
    s_checked, s_errors, _, _ = ref.compare_suivi(path)
    if errors or s_errors:
        raise SystemExit(f"{meta['id']}: reference mismatch {errors[:3]} {s_errors[:3]}")

    inputs = {
        "budget": {
            "incomeLines": lines(b, range(4, 7)),
            "fixedCostLines": lines(b, range(10, 18)),
            "variableExpenseLines": lines(b, range(21, 25)),
            "startMonth": month_of(b["B28"].value),
            "movingGoal": num(b["B29"].value),
            "movingDeadlineMonth": month_of(b["B30"].value),
            "movingAlreadySaved": num(b["B31"].value),
            "emergencyTarget": num(b["B32"].value),
            "emergencyExisting": num(b["B34"].value),
            "riskFreeRate": num(b["B35"].value),
            "earlyRepaymentPct": num(b["B36"].value),
        },
        # null = blank cell (template rows are all filled; no_loans has 6 blank slots)
        "loans": [
            {
                "slot": i + 1,
                "name": c[f"A{r}"].value,
                "type": c[f"B{r}"].value,
                "principal": num(c[f"C{r}"].value),
                "apr": num(c[f"D{r}"].value),
                "monthlyPayment": num(c[f"E{r}"].value),
            }
            for i, r in enumerate(ref.LOAN_ROWS)
        ],
        "actuals": [],
    }
    for r in ref.SUIVI_ROWS:
        raw = [su.cell(r, col).value for col in range(2, 13)]
        if any(isinstance(x, (int, float)) for x in raw):
            inputs["actuals"].append(
                {
                    "month": month_of(su[f"A{r}"].value),
                    "income": num(raw[0]),
                    "expenses": num(raw[1]),
                    "movingSavings": num(raw[2]),
                    "emergencySavings": num(raw[3]),
                    "freeSavings": num(raw[4]),
                    "loanBalances": [num(x) for x in raw[5:11]],  # by slot 1..6, null = blank
                }
            )

    # per-loan aggregates of Excel's own Calcul values (replace the dropped 'remaining months' columns)
    def payoff(slot_idx, col_off):
        start = ref.CALC_BLOCKS[slot_idx]
        if (num(k[f"{start}3"].value) or 0) <= 0:
            return None
        for m in range(ref.HORIZON):
            if (num(k[f"{ref.col_offset(start, col_off)}{3 + m}"].value) or 0) <= 0.01:
                return month_of(p[f"B{5 + m}"].value)
        return None

    def loan_sum(slot_idx, col_off):
        start = ref.CALC_BLOCKS[slot_idx]
        return num(sum(num(k[f"{ref.col_offset(start, col_off)}{3 + m}"].value) or 0 for m in range(ref.HORIZON)))

    credits = [
        {
            "slot": i + 1,
            "payoffMonthWithPlan": payoff(i, 5),
            "payoffMonthWithoutPlan": payoff(i, 6),
            "interestWithPlan": loan_sum(i, 1),
            "interestWithoutPlan": loan_sum(i, 7),
            "earlyRepaymentWorthIt": val(c[f"J{r}"].value),
            "priority": val(c[f"K{r}"].value),
            "advice": val(c[f"L{r}"].value),
        }
        for i, r in enumerate(ref.LOAN_ROWS)
    ]
    credits_totals = {
        "principal": num(c["C11"].value),
        "weightedApr": num(c["D11"].value),
        "monthlyPayments": num(c["E11"].value),
    }

    plan_cols = {
        "month": "A", "yearMonth": "B", "income": "C", "expenses": "D", "loanPayments": "E", "available": "F",
        "toMoving": "G", "movingCumulative": "H", "toEmergency": "I", "emergencyCumulative": "J",
        "remainder": "K", "toEarlyRepayment": "L", "unusedEarlyRepayment": "M", "toFreeSavings": "N",
        "freeSavingsCumulative": "O", "remainingDebt": "P",
    }
    plan = []
    for i in range(ref.HORIZON):
        r = 5 + i
        row = {key: val(p[f"{col}{r}"].value) for key, col in plan_cols.items()}
        row["month"] = int(row["month"])
        row["negativeBudget"] = p[f"Q{r}"].value == "Budget négatif !"
        row["movingReached"] = p[f"R{r}"].value == 1
        row["emergencyReached"] = p[f"S{r}"].value == 1
        row["debtFree"] = p[f"T{r}"].value == 1
        plan.append(row)

    sub = ["startBalance", "interest", "paymentPaid", "balanceAfterPayment", "earlyRepayment", "endBalance",
           "baselineEndBalance", "baselineInterest"]
    calcul = []
    for i in range(ref.HORIZON):
        r = 3 + i
        loans = []
        for start in ref.CALC_BLOCKS:
            loans.append([num(k[f"{ref.col_offset(start, si)}{r}"].value) for si in range(len(sub))])
        calcul.append(
            {
                "month": i + 1,
                "loans": loans,  # index = slot - 1; each row follows calculLoanColumns
                "totalPayments": num(k[f"AX{r}"].value),
                "totalEarlyRepayment": num(k[f"AY{r}"].value),
                "totalEndDebt": num(k[f"AZ{r}"].value),
                "totalInterest": num(k[f"BA{r}"].value),
                "totalBaselineInterest": num(k[f"BB{r}"].value),
            }
        )

    syn_map = {
        "monthlyIncome": "B4", "monthlyExpenses": "B5", "monthlyLoanPayments": "B6", "margin": "B7",
        "debtRatio": "B8", "debtAlert": "B9", "totalPrincipal": "B10", "weightedApr": "B11",
        "movingGoal": "B14", "movingMonthlyNeeded": "B15", "movingAmountAtDeadline": "B16",
        "movingReachedDate": "B17", "movingStatus": "B18", "emergencyTarget": "B21",
        "emergencyReachedDate": "B22", "debtFreeDate": "B25", "interestWithoutPlan": "B26",
        "interestWithPlan": "B27", "interestSaved": "B28", "freeSavingsAt12": "B31",
        "emergencyFundAt12": "B32", "remainingDebtAt12": "B33", "negativeBudgetMonths": "B34",
        "latestActualMonth": "B37", "latestDebtGap": "B38", "latestSavingsGap": "B39", "latestStatus": "B40",
    }
    synthese = {key: val(s[coord].value) for key, coord in syn_map.items()}

    suivi = []
    for r in ref.SUIVI_ROWS:
        if su[f"V{r}"].value in (0, None):
            continue
        suivi.append(
            {
                "month": month_of(su[f"A{r}"].value),
                "actualDebt": num(su[f"M{r}"].value),
                "plannedDebt": num(su[f"N{r}"].value),
                "debtGap": num(su[f"O{r}"].value),
                "actualSavings": num(su[f"P{r}"].value),
                "plannedSavings": num(su[f"Q{r}"].value),
                "savingsGap": num(su[f"R{r}"].value),
                "movingGoalPct": num(su[f"S{r}"].value),
                "debtRepaidPct": num(su[f"T{r}"].value),
                "status": su[f"U{r}"].value,
            }
        )

    return {
        "id": meta["id"],
        "description": meta["description"],
        "inputs": inputs,
        "expected": {
            "credits": credits,
            "creditsTotals": credits_totals,
            "plan": plan,
            "calcul": calcul,
            "synthese": synthese,
            "suivi": suivi,
        },
        "_verifiedCells": checked + s_checked,
    }


def main():
    src_dir = sys.argv[1]
    with open(os.path.join(ROOT, "scripts", "scenarios.json"), encoding="utf-8") as f:
        spec = json.load(f)
    scenarios = []
    for meta in spec["scenarios"]:
        sc = export_scenario(os.path.join(src_dir, meta["id"] + ".xlsx"), meta)
        print(f"{meta['id']}: ok ({sc.pop('_verifiedCells')} cells cross-checked)")
        scenarios.append(sc)
    golden = {
        "version": 1,
        "source": "docs/plan_financier.template.xlsx (anonymized) + xx.xx ROUND patches (scripts/scenarios.json), recalculated by desktop Excel per scenario",
        "horizonMonths": ref.HORIZON,
        "calculLoanColumns": ["startBalance", "interest", "paymentPaid", "balanceAfterPayment", "earlyRepayment", "endBalance", "baselineEndBalance", "baselineInterest"],
        "tolerance": 0.01,
        "notes": [
            "Expected values are the spreadsheet's own cached values (xx.xx-patched), rounded to 9 decimals.",
            "Scenario 'base' is the anonymized template; others override a few inputs (scripts/scenarios.json). ALL DATA IS INVENTED.",
            "Months are 'YYYY-MM'. Text KPIs keep the exact French strings (emojis included).",
            "credits[].payoffMonth*/interest* are sums/first-hits over Excel's Calcul values (SPEC §3).",
            "Loan arrays are indexed by slot (1..6); null = blank cell in the spreadsheet.",
        ],
        "scenarios": scenarios,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(golden, f, ensure_ascii=False, separators=(",", ":"))
    print(f"wrote {OUT} ({os.path.getsize(OUT) / 1e6:.2f} MB)")


if __name__ == "__main__":
    main()
