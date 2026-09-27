"""One-off: semantic diff of two workbooks (ignores quoting / int-vs-float noise). Not shipped."""
import sys
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

def norm(v):
    if isinstance(v, ArrayFormula): v = v.text
    if isinstance(v, float) and v.is_integer(): v = int(v)
    if isinstance(v, str): v = v.replace("'Crédits'!", "Crédits!").replace("'Suivi réel'!", "Suivi réel!")
    return v

def key(k):
    col, row = coordinate_from_string(k); return (row, column_index_from_string(col))

a = openpyxl.load_workbook(sys.argv[1]); b = openpyxl.load_workbook(sys.argv[2])
limit = int(sys.argv[3]) if len(sys.argv) > 3 else 400
for name in b.sheetnames:
    if name not in a.sheetnames:
        print(f"## NEW SHEET {name}"); continue
    wa, wb_ = a[name], b[name]
    cells = {c.coordinate for ws in (wa, wb_) for row in ws.iter_rows() for c in row if c.value is not None}
    diffs = [(k, norm(wa[k].value), norm(wb_[k].value)) for k in sorted(cells, key=key)]
    diffs = [d for d in diffs if d[1] != d[2]]
    print(f"## {name}: {len(diffs)} diffs")
    for k, x, y in diffs[:limit]:
        print(f"{k}: {x!r}\n   -> {y!r}")
